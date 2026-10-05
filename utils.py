import html
import re
from io import BytesIO

def sanitize_text(text: str) -> str:
    text = text.replace("“", '"').replace("”", '"')
    text = text.replace("‘", "'").replace("’", "'")
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", text)
    return text.strip()

def format_html_preview(text: str) -> str:
    safe = html.escape(sanitize_text(text))
    safe = safe.replace("\n", "<br>")
    return f"""
    <div style="
        background:#111827;
        color:#f9fafb;
        padding:24px;
        border-radius:12px;
        line-height:1.7;
        font-family:Georgia, serif;
        max-height:600px;
        overflow-y:auto;">
        {safe}
    </div>
    """

def format_docx(text: str, doc_type: str) -> bytes:
    from docx import Document
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(doc_type.upper())
    run.bold = True
    run.font.size = Pt(18)

    notice = doc.add_paragraph("AI-GENERATED DRAFT — LEGAL REVIEW RECOMMENDED")
    notice.alignment = WD_ALIGN_PARAGRAPH.CENTER
    notice.runs[0].italic = True

    for line in sanitize_text(text).splitlines():
        p = doc.add_paragraph()
        r = p.add_run(line)
        r.font.name = "Times New Roman"
        r.font.size = Pt(11)

    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("LegalEase — AI Legal Document Generator")

    output = BytesIO()
    doc.save(output)
    return output.getvalue()

def format_pdf(text: str, doc_type: str) -> bytes:
    from fpdf import FPDF

    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=18)

    def header():
        pdf.set_font("Helvetica", "B", 14)
        pdf.cell(0, 10, "LegalEase", align="C")
        pdf.ln(8)

    def footer():
        pdf.set_y(-15)
        pdf.set_font("Helvetica", size=8)
        pdf.cell(0, 10, "LegalEase - AI-generated draft | Legal review recommended", align="C")

    pdf.add_page()
    pdf.header = header
    pdf.footer = footer

    header()
    pdf.set_font("Helvetica", "B", 16)
    pdf.multi_cell(0, 9, sanitize_text(doc_type).upper(), align="C")
    pdf.ln(4)

    pdf.set_font("Helvetica", size=10)
    for line in sanitize_text(text).splitlines():
        if not line.strip():
            pdf.ln(3)
        else:
            pdf.multi_cell(0, 6, line)

    footer()
    return bytes(pdf.output())

def terms_table_data(terms: str):
    return [item.strip() for item in terms.split(";") if item.strip()]
